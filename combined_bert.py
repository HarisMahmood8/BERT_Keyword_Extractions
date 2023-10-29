import docx
from transformers import BertTokenizer, BertForSequenceClassification
import torch
import openpyxl
# python program that reads the .xlsx file containing keywords & topics
from test_keyword_read import read_keyword_topics, export_to_excel
# used to get rid of stopwords when checking for keywords
import nltk
from nltk.corpus import stopwords
from nltk import word_tokenize

# Load the BERT model
model = BertForSequenceClassification.from_pretrained("bert-base-uncased")
tokenizer = BertTokenizer.from_pretrained("bert-base-uncased")

# Load the DOCX file
quarter = "Q1"
doc = docx.Document(f"Equifax_{quarter}.docx")
output_file = f"sentiment_results_{quarter}.xlsx"

# Download stopwords
nltk.download('punkt')
nltk.download('stopwords')
stop_words = set(stopwords.words('english'))


# Define a function for sentiment analysis
def analyze_sentiment(text):
    inputs = tokenizer(text, return_tensors="pt", truncation=True, padding=True)
    outputs = model(**inputs)
    positive_logit = outputs.logits[0][1]
    negative_logit = outputs.logits[0][0]
    sentiment_score = positive_logit - negative_logit
    return sentiment_score


# For all found keywords, return a list of categories they fall under
def get_categories(keywords_list, categories):
    category_list = set()
    for keyword in keywords_list:
        for category in categories:
            if keyword in categories.get(category):
                category_list.add(category)
    return category_list


# Get dictionary of keywords and topics
workbook = openpyxl.load_workbook("keywords_topics.xlsx")
categories_dict = read_keyword_topics(workbook["Key Words_Topics"])
keywords = [item for sublist in categories_dict.values() for item in sublist]

# Create a list storing results, which will be used to create the excel output
result_list = []

# Process the document for sentiment
for paragraph in doc.paragraphs:
    text = paragraph.text
    sentiment = analyze_sentiment(text)
    text_sentences = text.split(". ")

    for sentence in text_sentences:
        text_words = sentence.split()

        found_keywords = [keyterm for keyterm in keywords if any(keyword in text_words and word_tokenize(keyword)[0].isalpha() and keyword.lower() not in stop_words for keyword in keyterm.split())]

        if found_keywords:
            categories_list = get_categories(found_keywords, categories_dict)
            result_list.append([sentence, float(sentiment), ', '.join(found_keywords), ', '.join(categories_list)])

# Output results to excel sheet
export_to_excel(result_list, output_file)

workbook.close()

print(f"Sentiment analysis and keyword extraction completed. Results saved in '{output_file}'.")
